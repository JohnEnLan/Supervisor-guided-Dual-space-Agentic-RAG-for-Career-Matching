import json
import os
from pathlib import Path

import pytest
from docx import Document

os.environ.setdefault("DATABASE_URL", "postgresql://test:test@localhost:5432/test")
os.environ.setdefault("DEEPSEEK_API_KEY", "sk-test")
os.environ.setdefault("QWEN_API_KEY", "sk-test")


def test_docx_extraction_preserves_paragraph_and_table_order(tmp_path):
    from app.normalization import resume_intake as intake

    resume_path = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("Profile first")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "University of Birmingham"
    table.cell(0, 1).text = "MSc Computer Science"
    document.add_paragraph("Experience last")
    document.save(resume_path)

    text, page_count = intake.extract_resume_text(resume_path)

    assert page_count == 1
    assert text.splitlines() == [
        "Profile first",
        "University of Birmingham",
        "MSc Computer Science",
        "Experience last",
    ]


def test_normalization_prompt_contains_only_retained_span_text():
    from app.normalization import resume_intake as intake

    prompt = intake._build_user_prompt(
        "UNRETAINED RAW RESUME SECRET",
        [intake.EvidenceSpan(span_id="R001", text="Python data analysis")],
    )
    payload = json.loads(prompt)

    assert "UNRETAINED RAW RESUME SECRET" not in prompt
    assert payload == {
        "retained_resume_text": "[R001] Python data analysis",
        "evidence_spans": [
            {
                "span_id": "R001",
                "page": None,
                "text": "Python data analysis",
                "source": "resume",
            }
        ],
    }


@pytest.mark.asyncio
async def test_normalization_rejects_facts_without_retained_span_ids(monkeypatch):
    from app.normalization import resume_intake as intake

    async def fake_chat(_system, _user, **_kwargs):
        return json.dumps(
            {
                "education": [
                    {
                        "institution": "University of Birmingham",
                        "degree": "MSc",
                        "evidence_span_ids": ["R001", "MISSING"],
                    },
                    {
                        "institution": "Invented University",
                        "degree": "PhD",
                        "evidence_span_ids": ["MISSING"],
                    },
                    {
                        "institution": "Fabricated College",
                        "degree": "PhD",
                        "evidence_span_ids": ["R001"],
                    },
                ],
                "experience": [],
                "projects": [
                    {
                        "name": "Invented project",
                        "evidence_span_ids": [],
                    }
                ],
                "skills": [
                    {"skill": "Python", "evidence_span_ids": ["R001"]},
                    {"skill": "Go", "evidence_span_ids": ["MISSING"]},
                    {"skill": "Rust", "evidence_span_ids": ["R001"]},
                    "Rust",
                ],
                "resume_quality_issues": [
                    {
                        "issue": "Dense formatting",
                        "severity": "medium",
                        "evidence_span_ids": ["R001"],
                    },
                    {
                        "issue": "Invented issue",
                        "severity": "high",
                        "evidence_span_ids": ["MISSING"],
                    },
                ],
                "normalized_base_resume": "Python data analysis",
            }
        )

    monkeypatch.setattr(intake, "chat", fake_chat)
    spans = [
        intake.EvidenceSpan(
            span_id="R001",
            text="University of Birmingham MSc; Python data analysis",
        )
    ]

    result = await intake.normalize_resume_text("ignored raw text", spans)

    assert result.education == [
        {
            "institution": "University of Birmingham",
            "degree": "MSc",
            "evidence_span_ids": ["R001"],
        }
    ]
    assert result.projects == []
    assert result.skills == ["Python"]
    assert result.resume_quality_issues == [
        "unverified: medium: Dense formatting evidence=['R001']"
    ]


@pytest.mark.asyncio
async def test_normalization_keeps_paraphrased_real_items_as_unverified(monkeypatch):
    """真实经历被模型改写（日期重排、措辞润色）时不再整条丢弃：
    身份锚点（公司/职位/项目名）能逐字对上证据 → 保留并标记 unverified；
    身份锚点全对不上的条目仍按编造拒绝。"""
    from app.normalization import resume_intake as intake

    async def fake_chat(_system, _user, **_kwargs):
        return json.dumps(
            {
                "education": [],
                "experience": [
                    {
                        "organization": "杭州云帆科技有限公司",
                        "title": "数据分析实习",
                        "dates": "2024.06-2024.08",
                        "location": "",
                        "responsibilities": ["搭建销售数据看板并输出周报"],
                        "achievements": [],
                        "technologies": [],
                        "evidence_span_ids": ["R003"],
                    },
                    {
                        "organization": "编造集团",
                        "title": "首席科学家",
                        "dates": "2020-2024",
                        "responsibilities": ["领导百人团队"],
                        "evidence_span_ids": ["R003"],
                    },
                ],
                "projects": [],
                "skills": [],
                "resume_quality_issues": [],
                "normalized_base_resume": "数据分析实习",
            }
        )

    monkeypatch.setattr(intake, "chat", fake_chat)
    spans = [
        intake.EvidenceSpan(
            span_id="R003",
            text="杭州云帆科技有限公司 数据分析实习 2024年6月至8月，负责销售数据看板的搭建与每周汇报",
        )
    ]

    result = await intake.normalize_resume_text("ignored", spans)

    assert len(result.experience) == 1
    kept = result.experience[0]
    assert kept["organization"] == "杭州云帆科技有限公司"
    assert kept["verification_status"] == "unverified"
    assert kept["evidence_span_ids"] == ["R003"]


@pytest.mark.asyncio
async def test_intake_resume_save_to_db_does_not_close_global_pool(monkeypatch):
    from app.normalization import resume_intake as intake
    from app.state.schema import ResumeState

    calls = []

    async def fake_normalize_resume_text(raw_text, evidence_spans):
        return ResumeState(
            normalized_base_resume="normalized resume",
            original_evidence_spans=[span.model_dump() for span in evidence_spans],
        )

    async def fake_save_state(state, status):
        calls.append(("save_state", state.session_id, status))

    async def fake_close_pool():
        calls.append(("close_pool",))

    monkeypatch.setattr(intake, "extract_resume_text", lambda path: ("raw resume", 1))
    monkeypatch.setattr(
        intake,
        "build_evidence_spans",
        lambda raw_text: [intake.EvidenceSpan(span_id="R001", text=raw_text)],
    )
    monkeypatch.setattr(intake, "normalize_resume_text", fake_normalize_resume_text)
    monkeypatch.setattr(intake, "save_state", fake_save_state)
    monkeypatch.setattr(intake, "close_pool", fake_close_pool)

    result = await intake.intake_resume(
        Path("resume.pdf"),
        session_id="session-1",
        user_id="user-1",
        save_to_db=True,
    )

    assert result.state.session_id == "session-1"
    assert ("save_state", "session-1", "resume_normalized") in calls
    assert ("close_pool",) not in calls


@pytest.mark.asyncio
async def test_intake_offloads_blocking_resume_parser_with_to_thread(
    monkeypatch,
) -> None:
    from app.normalization import resume_intake as intake
    from app.state.schema import ResumeState

    offloaded = []

    def parse(path):
        assert path == Path("resume.pdf")
        return "raw resume", 1

    async def to_thread(function, *args):
        offloaded.append((function, args))
        return function(*args)

    async def normalize(raw_text, evidence_spans):
        return ResumeState(
            normalized_base_resume=raw_text,
            original_evidence_spans=[span.model_dump() for span in evidence_spans],
        )

    monkeypatch.setattr(intake, "extract_resume_text", parse)
    monkeypatch.setattr(intake.asyncio, "to_thread", to_thread)
    monkeypatch.setattr(intake, "normalize_resume_text", normalize)

    await intake.intake_resume(
        Path("resume.pdf"),
        session_id="session-1",
        user_id="user-1",
    )

    assert offloaded == [(parse, (Path("resume.pdf"),))]


def test_clarification_trigger_normalizes_prioritizes_deduplicates_and_caps() -> None:
    from app.normalization import resume_intake as intake

    quality_issues, targets = intake.build_clarification_targets(
        quality_issues=[
            {
                "issue": "缺少量化成果",
                "severity": "HIGH",
                "field_path": "experience[0]",
                "evidence_span_ids": ["R001"],
            },
            {
                "issue": "教育经历说明不够清楚",
                "severity": "unexpected",
                "field_path": "education[0]",
                "evidence_span_ids": ["R002"],
            },
            {
                "issue": "轻微排版问题",
                "severity": "low",
                "field_path": "resume_quality_issues[2]",
                "evidence_span_ids": ["R003"],
            },
        ],
        experience=[
            {
                "organization": "Acme",
                "responsibilities": ["维护接口"],
                "achievements": [],
                "evidence_span_ids": ["R001"],
            },
            {
                "organization": "Beta",
                "responsibilities": ["负责数据管道的设计、开发和日常稳定性维护"],
                "achievements": ["将批处理时间从两小时缩短到四十分钟"],
                "evidence_span_ids": ["R004"],
            },
        ],
        projects=[
            {
                "name": "Career RAG",
                "summary": "为求职者提供岗位匹配建议",
                "technologies": [],
                "evidence_span_ids": ["R005"],
            }
        ],
        max_targets=2,
    )

    assert [item["severity"] for item in quality_issues] == [
        "high",
        "medium",
        "low",
    ]
    assert [item["target_ref"] for item in targets] == [
        "experience[0]",
        "education[0]",
    ]
    assert [item["severity"] for item in targets] == ["high", "medium"]
    assert all(item["status"] == "open" for item in targets)
    assert len({item["field_path"] for item in targets}) == len(targets)


def test_clarification_trigger_detects_missing_descriptions_and_project_skills() -> None:
    from app.normalization import resume_intake as intake

    _quality_issues, targets = intake.build_clarification_targets(
        quality_issues=[],
        experience=[
            {
                "organization": "Acme",
                "responsibilities": [],
                "achievements": [],
                "evidence_span_ids": ["R001"],
            }
        ],
        projects=[
            {
                "name": "Portfolio",
                "summary": "个人作品集",
                "technologies": [],
                "evidence_span_ids": ["R002"],
            }
        ],
        max_targets=5,
    )

    assert [(item["target_ref"], item["severity"]) for item in targets] == [
        ("experience[0]", "medium"),
        ("projects[0]", "medium"),
    ]
    assert "时间" not in " ".join(item["issue"] for item in targets)


@pytest.mark.asyncio
@pytest.mark.parametrize("enabled", [False, True])
async def test_normalization_writes_feature_a_fields_only_when_enabled(
    monkeypatch,
    enabled,
) -> None:
    from app.normalization import resume_intake as intake

    calls = 0

    async def fake_chat(_system, _user, **_kwargs):
        nonlocal calls
        calls += 1
        return json.dumps(
            {
                "education": [],
                "experience": [],
                "projects": [],
                "skills": [],
                "resume_quality_issues": [
                    {
                        "issue": "缺少职责说明",
                        "severity": "high",
                        "field_path": "experience[0]",
                        "evidence_span_ids": ["R001"],
                    }
                ],
                "normalized_base_resume": "Acme",
            },
            ensure_ascii=False,
        )

    monkeypatch.setattr(intake, "chat", fake_chat)
    monkeypatch.setattr(intake.settings, "resume_clarify_enabled", enabled)
    monkeypatch.setattr(intake.settings, "resume_clarify_max", 2)

    result = await intake.normalize_resume_text(
        "ignored",
        [intake.EvidenceSpan(span_id="R001", text="Acme")],
    )

    assert calls == 1
    if enabled:
        assert result.quality_issues_struct[0]["severity"] == "high"
        assert result.clarification_targets[0]["target_ref"] == "experience[0]"
    else:
        assert result.quality_issues_struct == []
        assert result.clarification_targets == []


@pytest.mark.asyncio
async def test_switch_off_field_path_metadata_preserves_baseline_verification(
    monkeypatch,
) -> None:
    from app.normalization import resume_intake as intake

    captured_system_prompts: list[str] = []

    async def fake_chat(system, _user, **_kwargs):
        captured_system_prompts.append(system)
        return json.dumps(
            {
                "education": [],
                "experience": [],
                "projects": [],
                "skills": [],
                "resume_quality_issues": [
                    {
                        "issue": "Acme",
                        "severity": "medium",
                        "field_path": "experience[0]",
                        "target_ref": "experience[0]",
                        "evidence_span_ids": ["R001"],
                    }
                ],
                "normalized_base_resume": "Acme",
            }
        )

    monkeypatch.setattr(intake, "chat", fake_chat)
    monkeypatch.setattr(intake.settings, "resume_clarify_enabled", False)

    result = await intake.normalize_resume_text(
        "ignored",
        [intake.EvidenceSpan(span_id="R001", text="Acme")],
    )

    assert "field_path" not in captured_system_prompts[0]
    assert result.resume_quality_issues == ["medium: Acme evidence=['R001']"]
