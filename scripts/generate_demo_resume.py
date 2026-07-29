"""Generate an upload-ready demo resume aligned with the bundled job dataset."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("data/resumes/demo_software_engineer_resume.docx")
FONT = "Arial"
NAVY = RGBColor(31, 78, 121)
DARK = RGBColor(32, 32, 32)
MUTED = RGBColor(90, 90, 90)


def set_font(run, *, size: float, bold: bool = False, color: RGBColor = DARK) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_bottom_border(paragraph, color: str = "1F4E79", size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    set_bottom_border(paragraph, size="5")
    set_font(paragraph.add_run(text.upper()), size=10.5, bold=True, color=NAVY)


def add_role(
    doc: Document,
    *,
    title: str,
    organization: str,
    dates: str,
    bullets: list[str],
) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(2)
    heading.paragraph_format.space_after = Pt(1)
    heading.paragraph_format.keep_with_next = True
    set_font(heading.add_run(f"{title} | {organization}"), size=10.2, bold=True)
    date_run = heading.add_run(f"  |  {dates}")
    set_font(date_run, size=9.5, color=MUTED)

    for text in bullets:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        paragraph.paragraph_format.space_after = Pt(1.5)
        paragraph.paragraph_format.line_spacing = 1.0
        set_font(paragraph.add_run(text), size=9.3)


def build_resume() -> None:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    set_font(name.add_run("ALEX CHEN"), size=20, bold=True, color=NAVY)

    target = doc.add_paragraph()
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.paragraph_format.space_after = Pt(1)
    set_font(
        target.add_run("SOFTWARE ENGINEER | PYTHON BACKEND & API DEVELOPMENT"),
        size=10.5,
        bold=True,
    )

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(5)
    set_font(
        contact.add_run(
            "Demo Candidate  |  alex.chen.demo@example.com  |  +1 555 010 2026  |  Open to Remote / Relocation"
        ),
        size=8.8,
        color=MUTED,
    )

    add_section_heading(doc, "Professional Summary")
    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(3)
    summary.paragraph_format.line_spacing = 1.05
    set_font(
        summary.add_run(
            "Software engineer with hands-on experience building Python backend services, REST APIs, "
            "PostgreSQL data workflows, automated tests, and cloud deployment pipelines. Experienced "
            "in API debugging, production support, root-cause analysis, Docker, AWS, Git, and CI/CD. "
            "Seeking Software Engineer, Backend Engineer, or Software Support Specialist opportunities."
        ),
        size=9.4,
    )

    add_section_heading(doc, "Technical Skills")
    skills = doc.add_paragraph()
    skills.paragraph_format.space_after = Pt(3)
    skills.paragraph_format.line_spacing = 1.05
    set_font(skills.add_run("Languages: "), size=9.3, bold=True)
    set_font(skills.add_run("Python, SQL, JavaScript, PHP"), size=9.3)
    set_font(skills.add_run("  |  Backend: "), size=9.3, bold=True)
    set_font(skills.add_run("FastAPI, REST APIs, async programming, JSON"), size=9.3)
    set_font(skills.add_run("\nData & Cloud: "), size=9.3, bold=True)
    set_font(skills.add_run("PostgreSQL, MySQL, Redis, AWS, Docker"), size=9.3)
    set_font(skills.add_run("  |  Engineering: "), size=9.3, bold=True)
    set_font(skills.add_run("Git, GitHub Actions, CI/CD, pytest, debugging, Linux"), size=9.3)

    add_section_heading(doc, "Experience")
    add_role(
        doc,
        title="Junior Software Engineer",
        organization="Northstar Digital Labs",
        dates="Jun 2024 - Present",
        bullets=[
            "Built and maintained asynchronous Python and FastAPI backend services with REST API endpoints, input validation, and PostgreSQL persistence.",
            "Diagnosed API failures from logs and traces, reproduced defects, implemented fixes, and added pytest regression tests to prevent repeat incidents.",
            "Containerized services with Docker and supported AWS cloud deployment through a GitHub Actions CI/CD pipeline.",
            "Optimized SQL queries and API response handling for reliable data access, monitoring, and production support.",
        ],
    )
    add_role(
        doc,
        title="Software Support Intern",
        organization="BluePeak Systems",
        dates="Jan 2024 - May 2024",
        bullets=[
            "Worked in a Help Desk environment investigating web application issues, testing REST APIs, querying PostgreSQL and MySQL records, and documenting root causes.",
            "Created Python utilities for log analysis and repeatable troubleshooting, reducing manual diagnostic steps.",
            "Collaborated through Git-based code review and supported JavaScript and PHP web development fixes with clear technical notes.",
        ],
    )

    add_section_heading(doc, "Projects")
    add_role(
        doc,
        title="Cloud-Deployed Job Matching API",
        organization="Independent Project",
        dates="2024",
        bullets=[
            "Developed a FastAPI service for job search and candidate matching using Python, PostgreSQL, SQL filtering, and ranked API responses.",
            "Added automated tests, Docker packaging, environment-based configuration, structured logging, and an AWS deployment workflow.",
        ],
    )

    add_section_heading(doc, "Education")
    education = doc.add_paragraph()
    education.paragraph_format.space_after = Pt(0)
    set_font(
        education.add_run("BSc Computer Science | Riverside University | 2020 - 2024"),
        size=9.6,
        bold=True,
    )
    detail = doc.add_paragraph()
    detail.paragraph_format.space_after = Pt(0)
    set_font(
        detail.add_run(
            "Relevant study: Software Engineering, Databases, Cloud Computing, Data Structures, Web Development"
        ),
        size=9.2,
    )

    core = doc.core_properties
    core.title = "Demo Software Engineer Resume"
    core.subject = "Upload-ready demonstration resume for Career-RAG"
    core.author = "Career-RAG Demo"
    core.keywords = (
        "Python, Software Engineer, Software Support Specialist, FastAPI, REST API, "
        "PostgreSQL, MySQL, PHP, JavaScript, AWS, Docker, Debugging, Help Desk"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_resume()
