"""Render the V2 functionality/code guide Markdown into a styled DOCX."""

from __future__ import annotations

import argparse
import os
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = REPO_ROOT / "docs" / "project_functionality_and_code_guide.md"
DEFAULT_OUTPUT = REPO_ROOT / "docs" / "功能与代码详解_V2.docx"
INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
ORDERED_LIST_PATTERN = re.compile(r"^\s*\d+\.\s+(.+)$")
UNORDERED_LIST_PATTERN = re.compile(r"^\s*[-*+]\s+(.+)$")
TABLE_DIVIDER_PATTERN = re.compile(r"^:?-{3,}:?$")


@dataclass
class RenderStats:
    heading_1: int = 0
    heading_2: int = 0
    heading_3: int = 0
    tables: int = 0
    code_blocks: int = 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Render the Career-RAG V2 Markdown guide to DOCX."
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=DEFAULT_SOURCE,
        help=f"Markdown source (default: {DEFAULT_SOURCE})",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"DOCX output (default: {DEFAULT_OUTPUT})",
    )
    return parser.parse_args()


def backup_path_for(output: Path) -> Path:
    return output.with_name(f"{output.stem}_prev{output.suffix}")


def preserve_previous_output(output: Path) -> Path | None:
    """Create the one-time _prev backup and return the style template path."""
    backup = backup_path_for(output)
    if output.exists() and not backup.exists():
        shutil.copy2(output, backup)
        print(f"Backed up previous DOCX: {backup}")
    if backup.exists():
        return backup
    return None


def clear_document_body(document: DocumentObject) -> None:
    """Remove template body content while preserving its section properties."""
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def ensure_styles(document: DocumentObject) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for style_name, size, color in (
        ("Title", 24, "1F4E79"),
        ("Subtitle", 11, "5B6573"),
        ("Heading 1", 17, "1F4E79"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 11, "3F6B9A"),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.base_style = styles["Normal"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.5)
    code_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.08)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(7)

    if "Guide TOC" not in styles:
        toc_style = styles.add_style("Guide TOC", WD_STYLE_TYPE.PARAGRAPH)
    else:
        toc_style = styles["Guide TOC"]
    toc_style.base_style = styles["Normal"]
    toc_style.font.name = "Calibri"
    toc_style.font.size = Pt(10)
    toc_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    toc_style.paragraph_format.space_after = Pt(2)


def shade_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_inline_markdown(paragraph, text: str) -> None:
    """Render the small inline Markdown subset used by the guide."""
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", 9)
            run.font.color.rgb = RGBColor(0x7A, 0x21, 0x21)
        else:
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            run.underline = True
            paragraph.add_run(f" ({target})")
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or "|" not in lines[index]:
        return False
    divider = split_table_row(lines[index + 1])
    return bool(divider) and all(TABLE_DIVIDER_PATTERN.match(cell) for cell in divider)


def collect_paragraph(lines: list[str], start: int) -> tuple[str, int]:
    parts: list[str] = []
    index = start
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            break
        if index != start and (
            HEADING_PATTERN.match(line)
            or line.startswith("```")
            or line.startswith(">")
            or ORDERED_LIST_PATTERN.match(line)
            or UNORDERED_LIST_PATTERN.match(line)
            or is_table_start(lines, index)
        ):
            break
        parts.append(stripped)
        index += 1
    return " ".join(parts), index


def extract_title(lines: list[str]) -> tuple[str, str, list[str]]:
    title = "Career-RAG 功能与代码详解"
    subtitle = "Supervisor-guided Dual-space Agentic RAG for Career Matching"
    body = list(lines)
    if body and body[0].startswith("# "):
        title = body.pop(0)[2:].strip()
    while body and not body[0].strip():
        body.pop(0)
    if body and body[0].startswith(">"):
        subtitle = body.pop(0)[1:].strip()
    return title, subtitle, body


def extract_headings(lines: list[str]) -> list[tuple[int, str]]:
    headings: list[tuple[int, str]] = []
    in_code = False
    for line in lines:
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        match = HEADING_PATTERN.match(line)
        if match:
            headings.append((len(match.group(1)), match.group(2)))
    return headings


def add_title_page(document: DocumentObject, title: str, subtitle: str) -> None:
    document.add_paragraph()
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = 1
    add_inline_markdown(paragraph, title)
    subtitle_paragraph = document.add_paragraph(style="Subtitle")
    subtitle_paragraph.alignment = 1
    add_inline_markdown(subtitle_paragraph, subtitle)
    baseline = document.add_paragraph()
    baseline.alignment = 1
    run = baseline.add_run("冻结代码基线 2474e48 · 2026-08-07")
    run.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    document.add_page_break()


def add_static_toc(document: DocumentObject, headings: list[tuple[int, str]]) -> None:
    title = document.add_paragraph(style="Title")
    title.add_run("目录")
    note = document.add_paragraph()
    note.add_run("目录由 Markdown 标题自动生成。").italic = True
    for level, text in headings:
        if level > 3:
            continue
        paragraph = document.add_paragraph(style="Guide TOC")
        paragraph.paragraph_format.left_indent = Inches(0.22 * (level - 1))
        if level == 1:
            paragraph.runs.clear()
            run = paragraph.add_run(text)
            run.bold = True
        else:
            paragraph.add_run(text)
    document.add_page_break()


def add_table(document: DocumentObject, rows: list[list[str]]) -> None:
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.autofit = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    header_properties.append(header_flag)
    for row_index, values in enumerate(rows):
        for column_index in range(column_count):
            value = values[column_index] if column_index < len(values) else ""
            cell = table.cell(row_index, column_index)
            cell.text = ""
            add_inline_markdown(cell.paragraphs[0], value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
            if row_index == 0:
                shade_cell(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def render_body(
    document: DocumentObject, lines: list[str], stats: RenderStats
) -> None:
    index = 0
    first_part = True
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        heading = HEADING_PATTERN.match(line)
        if heading:
            level = min(len(heading.group(1)), 3)
            paragraph = document.add_heading(level=level)
            if level == 1 and not first_part:
                paragraph.paragraph_format.page_break_before = True
            add_inline_markdown(paragraph, heading.group(2))
            if level == 1:
                stats.heading_1 += 1
                first_part = False
            elif level == 2:
                stats.heading_2 += 1
            else:
                stats.heading_3 += 1
            index += 1
            continue

        if line.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            paragraph = document.add_paragraph(style="Code Block")
            shade_paragraph(paragraph, "F3F4F6")
            run = paragraph.add_run("\n".join(code_lines))
            set_run_font(run, "Consolas", 8.5)
            stats.code_blocks += 1
            continue

        if is_table_start(lines, index):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            stats.tables += 1
            continue

        ordered = ORDERED_LIST_PATTERN.match(line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, ordered.group(1))
            index += 1
            continue

        unordered = UNORDERED_LIST_PATTERN.match(line)
        if unordered:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, unordered.group(1))
            index += 1
            continue

        if line.startswith(">"):
            quote = document.add_paragraph(style="Intense Quote")
            add_inline_markdown(quote, line[1:].strip())
            index += 1
            continue

        if stripped == "---":
            index += 1
            continue

        text, next_index = collect_paragraph(lines, index)
        paragraph = document.add_paragraph()
        add_inline_markdown(paragraph, text)
        index = next_index


def build_document(source: Path, output: Path, template: Path | None) -> RenderStats:
    markdown = source.read_text(encoding="utf-8")
    title, subtitle, body_lines = extract_title(markdown.splitlines())
    document = Document(template) if template else Document()
    if template:
        clear_document_body(document)
    else:
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.945)
        section.right_margin = Inches(0.945)
        section.top_margin = Inches(0.866)
        section.bottom_margin = Inches(0.866)

    ensure_styles(document)
    document.core_properties.title = title
    document.core_properties.subject = "Career-RAG V2 functionality and code guide"
    document.core_properties.author = "Career-RAG project"

    add_title_page(document, title, subtitle)
    add_static_toc(document, extract_headings(body_lines))
    stats = RenderStats()
    render_body(document, body_lines, stats)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp")
    if temporary.exists():
        temporary.unlink()
    document.save(temporary)
    os.replace(temporary, output)
    return stats


def main() -> int:
    args = parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Markdown source not found: {source}")
    template = preserve_previous_output(output)
    stats = build_document(source, output, template)
    print(f"Rendered DOCX: {output}")
    print(
        "Structure: "
        f"Heading1={stats.heading_1}, Heading2={stats.heading_2}, "
        f"Heading3={stats.heading_3}, tables={stats.tables}, "
        f"code_blocks={stats.code_blocks}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
